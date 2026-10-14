"""Read RVTools exports and extract ESXi host hardware platforms (spec FR-1, FR-2)."""

from __future__ import annotations

import glob
import logging
import re
from collections import Counter
from dataclasses import dataclass, field
from importlib import resources
from pathlib import Path

import yaml

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".xlsx", ".xls")
VHOST_SHEET = "vhost"

# Logical field -> accepted header names (already normalized, see normalize_header).
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "host": ("host",),
    "vendor": ("vendor",),
    "model": ("model",),
    "datacenter": ("datacenter",),
    "cluster": ("cluster",),
    "cpu_model": ("cpu model",),
    "cpu_count": ("# cpu",),
    "cores_per_cpu": ("cores per cpu",),
    "memory": ("# memory",),
    "esx_version": ("esx version",),
    "serial": ("service tag", "serial number"),
}
REQUIRED_COLUMNS = ("host", "vendor", "model")


class RVToolsFileError(Exception):
    """The file cannot be read or is not an RVTools export."""


@dataclass
class Host:
    source_file: str
    host: str
    vendor: str
    model: str
    vendor_raw: str
    model_raw: str
    datacenter: str = ""
    cluster: str = ""
    cpu_model: str = ""
    cpu_count: str = ""
    cores_per_cpu: str = ""
    memory: str = ""
    esx_version: str = ""
    serial: str = ""

    @property
    def platform_key(self) -> tuple[str, str]:
        return (self.vendor, self.model)


@dataclass
class Platform:
    vendor: str
    model: str
    hosts: list[Host] = field(default_factory=list)

    @property
    def key(self) -> tuple[str, str]:
        return (self.vendor, self.model)

    @property
    def host_count(self) -> int:
        return len(self.hosts)

    @property
    def raw_names(self) -> list[str]:
        return sorted({f"{h.vendor_raw} {h.model_raw}".strip() for h in self.hosts})

    @property
    def files(self) -> dict[str, int]:
        """Host count per source file."""
        return dict(sorted(Counter(h.source_file for h in self.hosts).items()))

    @property
    def datacenters(self) -> dict[str, int]:
        """Host count per datacenter."""
        return dict(sorted(Counter(h.datacenter or "(none)" for h in self.hosts).items()))

    @property
    def lookup_possible(self) -> bool:
        return bool(self.vendor and self.model)


def collapse(value) -> str:
    """Stringify a cell, trim it and collapse repeated whitespace."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_header(value) -> str:
    return collapse(value).lower()


def _words(value: str) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", value.lower()).split())


class Normalizer:
    """Vendor and model normalization driven by vendor_aliases.yaml (FR-2.4, FR-2.5)."""

    def __init__(self, config: dict):
        self._aliases: list[tuple[str, str]] = []
        for canonical, aliases in (config.get("vendors") or {}).items():
            for alias in [canonical, *(aliases or [])]:
                self._aliases.append((_words(str(alias)), canonical))
        # Longest alias first so "hewlett packard enterprise" beats "hp".
        self._aliases.sort(key=lambda a: len(a[0]), reverse=True)
        self._rewrites = [
            (r.get("vendor"), re.compile(r["pattern"], re.IGNORECASE), r.get("replace", ""))
            for r in (config.get("model_rewrites") or [])
        ]

    @classmethod
    def from_file(cls, path: str | Path | None = None) -> "Normalizer":
        if path is None:
            text = resources.files("rvtools_hw_check").joinpath("config/vendor_aliases.yaml").read_text()
        else:
            text = Path(path).read_text()
        return cls(yaml.safe_load(text) or {})

    def vendor(self, raw) -> str:
        cleaned = collapse(raw)
        key = _words(cleaned)
        if not key:
            return ""
        for alias, canonical in self._aliases:
            if key == alias or key.startswith(alias + " "):
                return canonical
        return cleaned

    def model(self, vendor: str, raw) -> str:
        model = collapse(raw)
        for rule_vendor, pattern, replace in self._rewrites:
            if rule_vendor is None or rule_vendor == vendor:
                model = pattern.sub(replace, model)
        return collapse(model)


def expand_inputs(inputs: list[str], recursive: bool = False) -> tuple[list[Path], list[tuple[str, str]]]:
    """Resolve paths, globs and directories into RVTools files (FR-1.1).

    Directories contribute every .xlsx/.xls file in them; with `recursive`,
    their subfolders too. Returns (files, errors) where errors are (input, reason) pairs.
    """
    files: list[Path] = []
    errors: list[tuple[str, str]] = []
    for item in inputs:
        path = Path(item)
        if path.is_dir():
            entries = path.rglob("*") if recursive else path.iterdir()
            found = sorted(p for p in entries if _is_candidate(p))
            if not found:
                errors.append((item, "Directory contains no .xlsx or .xls files"))
            files.extend(found)
        elif path.is_file():
            files.append(path)
        else:
            matches = sorted(Path(m) for m in glob.glob(item))
            matches = [m for m in matches if _is_candidate(m)]
            if not matches:
                errors.append((item, "File not found"))
            files.extend(matches)

    unique: list[Path] = []
    seen: set[Path] = set()
    for f in files:
        resolved = f.resolve()
        if resolved not in seen:
            seen.add(resolved)
            unique.append(f)
    return unique, errors


def _is_candidate(path: Path) -> bool:
    # "~$" files are Excel lock files, not exports.
    return path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS and not path.name.startswith("~$")


def _read_sheet_rows(path: Path) -> list[list]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise RVToolsFileError(f"Unsupported file type '{path.suffix}' (expected .xlsx or .xls)")
    try:
        if suffix == ".xlsx":
            import openpyxl

            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                sheet_name = _find_sheet(workbook.sheetnames)
                return [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
            finally:
                workbook.close()
        import xlrd

        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            sheet = workbook.sheet_by_name(_find_sheet(workbook.sheet_names()))
            return [sheet.row_values(i) for i in range(sheet.nrows)]
        finally:
            workbook.release_resources()
    except RVToolsFileError:
        raise
    except Exception as exc:  # openpyxl/xlrd raise many unrelated types for bad files
        raise RVToolsFileError(f"Cannot read workbook: {exc}") from exc


def _find_sheet(names: list[str]) -> str:
    for name in names:
        if name.strip().lower() == VHOST_SHEET:
            return name
    raise RVToolsFileError("No 'vHost' sheet found; this does not look like an RVTools export")


def _map_columns(header: list) -> dict[str, int]:
    normalized = [normalize_header(h) for h in header]
    mapping: dict[str, int] = {}
    for field_name, aliases in COLUMN_ALIASES.items():
        for idx, name in enumerate(normalized):
            if name in aliases:
                mapping[field_name] = idx
                break
    missing = [c for c in REQUIRED_COLUMNS if c not in mapping]
    if missing:
        raise RVToolsFileError(
            "vHost sheet is missing required column(s): " + ", ".join(m.capitalize() for m in missing)
        )
    return mapping


def display_names(paths: list[Path]) -> dict[Path, str]:
    """Label each file by its name, or by its full path when two files share a name."""
    counts = Counter(p.name for p in paths)
    return {p: p.name if counts[p.name] == 1 else str(p) for p in paths}


def load_hosts(path: Path, normalizer: Normalizer, display_name: str | None = None) -> list[Host]:
    """Read the vHost sheet of one RVTools export (FR-2.1 - FR-2.5)."""
    rows = _read_sheet_rows(path)
    if not rows:
        raise RVToolsFileError("vHost sheet is empty")
    columns = _map_columns(rows[0])

    def cell(row: list, name: str) -> str:
        idx = columns.get(name)
        return collapse(row[idx]) if idx is not None and idx < len(row) else ""

    hosts: list[Host] = []
    for row in rows[1:]:
        host_name = cell(row, "host")
        if not host_name:
            continue
        vendor_raw = cell(row, "vendor")
        model_raw = cell(row, "model")
        vendor = normalizer.vendor(vendor_raw)
        hosts.append(
            Host(
                source_file=display_name or path.name,
                host=host_name,
                vendor=vendor,
                model=normalizer.model(vendor, model_raw),
                vendor_raw=vendor_raw,
                model_raw=model_raw,
                datacenter=cell(row, "datacenter"),
                cluster=cell(row, "cluster"),
                cpu_model=cell(row, "cpu_model"),
                cpu_count=cell(row, "cpu_count"),
                cores_per_cpu=cell(row, "cores_per_cpu"),
                memory=cell(row, "memory"),
                esx_version=cell(row, "esx_version"),
                serial=cell(row, "serial"),
            )
        )
    logger.info("Read %d host(s) from %s", len(hosts), display_name or path.name)
    return hosts


def group_platforms(hosts: list[Host]) -> list[Platform]:
    """Deduplicate hosts into unique (vendor, model) platforms across all files (FR-2.6)."""
    platforms: dict[tuple[str, str], Platform] = {}
    for host in hosts:
        platform = platforms.setdefault(host.platform_key, Platform(host.vendor, host.model))
        platform.hosts.append(host)
    return sorted(platforms.values(), key=lambda p: (p.vendor.lower(), p.model.lower()))
