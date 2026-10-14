"""Generate sample RVTools exports for tests and manual runs.

    python tests/fixtures/make_fixtures.py

writes rvtools_sample_a.xlsx (current RVTools layout) and rvtools_sample_b.xls
(older layout with differently-cased headers) next to this script.
"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).parent

# Current RVTools column order (subset), as exported to .xlsx.
HEADERS_A = [
    "Host", "Datacenter", "Cluster", "Config status", "CPU Model", "Speed", "# CPU", "Cores per CPU",
    "# Cores", "# Memory", "ESX Version", "Vendor", "Model", "Serial number", "Service tag",
]
ROWS_A = [
    ["esx01.lab.local", "DC1", "Prod", "green", "Intel(R) Xeon(R) Gold 6338 CPU @ 2.00GHz", 2000, 2, 32, 64, 524288,
     "VMware ESXi 7.0.3 build-21424296", "Dell Inc.", "PowerEdge R650", "ABC1234", "ABC1234"],
    ["esx02.lab.local", "DC1", "Prod", "green", "Intel(R) Xeon(R) Gold 6338 CPU @ 2.00GHz", 2000, 2, 32, 64, 524288,
     "VMware ESXi 7.0.3 build-21424296", "Dell Inc.", "PowerEdge  R650 ", "ABC1235", "ABC1235"],
    ["esx03.lab.local", "DC1", "Dev", "green", "Intel(R) Xeon(R) Gold 6230 CPU @ 2.10GHz", 2100, 2, 20, 40, 393216,
     "VMware ESXi 7.0.3 build-21424296", "HPE", "ProLiant DL380 Gen10", "CZJ0001", ""],
    ["esx04.lab.local", "DC2", "Edge", "green", "Intel(R) Xeon(R) Silver 4214 CPU @ 2.20GHz", 2200, 2, 12, 24, 262144,
     "VMware ESXi 6.7.0 build-17700523", "Cisco Systems Inc", "UCSC-C220-M5SX", "WZP0001", ""],
    [None, None, None, None, None, None, None, None, None, None, None, None, None, None, None],
]

# Older RVTools: fewer columns, different header spelling/case.
HEADERS_B = ["Host", "Datacenter", "Cluster", "CPU model", "# CPU", "Cores per CPU", "# Memory", "ESX Version", "  vendor ", "MODEL", "Service tag"]
ROWS_B = [
    ["esx10.corp.local", "DC-East", "Cluster-A", "Intel(R) Xeon(R) Gold 6338 CPU @ 2.00GHz", 2, 32, 524288,
     "VMware ESXi 6.7.0", "Dell Inc.", "PowerEdge R650", "XYZ0001"],
    ["esx11.corp.local", "DC-East", "Cluster-A", "Intel(R) Xeon(R) Gold 6130 CPU @ 2.10GHz", 2, 16, 393216,
     "VMware ESXi 6.7.0", "Lenovo", "ThinkSystem SR650 -[7X06CTO1WW]-", "J300001"],
    ["esx12.corp.local", "DC-East", "Cluster-B", "AMD EPYC 7302", 1, 16, 131072,
     "VMware ESXi 6.5.0", "Acme Corp", "Widget 9000", "ACME001"],
]


def write_xlsx(path: Path, headers: list, rows: list, sheet: str = "vHost") -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    info = workbook.active
    info.title = "vInfo"
    info.append(["VM", "Powerstate", "Host"])
    info.append(["vm01", "poweredOn", "esx01.lab.local"])
    vhost = workbook.create_sheet(sheet)
    vhost.append(headers)
    for row in rows:
        vhost.append(row)
    workbook.save(path)
    return path


def write_xls(path: Path, headers: list, rows: list, sheet: str = "vHost") -> Path:
    import xlwt

    workbook = xlwt.Workbook()
    workbook.add_sheet("vInfo").write(0, 0, "VM")
    vhost = workbook.add_sheet(sheet)
    for col, value in enumerate(headers):
        vhost.write(0, col, value)
    for r, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            if value is not None:
                vhost.write(r, col, value)
    workbook.save(str(path))
    return path


def make_samples(directory: Path) -> tuple[Path, Path]:
    directory.mkdir(parents=True, exist_ok=True)
    return (
        write_xlsx(directory / "rvtools_sample_a.xlsx", HEADERS_A, ROWS_A),
        write_xls(directory / "rvtools_sample_b.xls", HEADERS_B, ROWS_B),
    )


if __name__ == "__main__":
    for p in make_samples(HERE):
        print("wrote", p)
